"""Phase 6 — openpyxl-based Excel export. RBAC §2.2(f): "Excel is a plain
data export of the same filtered/sorted rows and visible columns... not a
formatted report" — deliberately no styling beyond a bold header row and
column widths, unlike the PDF path's letterhead treatment.
"""

from __future__ import annotations

from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

from app.services.pdf_engine import ColumnSpec

_FORMULA_TRIGGER_CHARS = ("=", "+", "-", "@", "\t", "\r")


def _sanitize_cell(v):
    """CWE-1236 (Formula/CSV injection): a spreadsheet app treats any cell
    beginning with =, +, -, @ (or a leading tab/CR, both used historically
    to smuggle a leading = past naive checks) as a live formula on open —
    confirmed live against this openpyxl version (a "=HYPERLINK(...)" cell
    value is auto-marked as a formula, data_type='f'). Every export row in
    this codebase can contain genuinely user-authored free text (a Vendor/
    Customer/Employee name, a Ticket subject, ...), so every string cell
    needs this check, not just ones that "look" attacker-supplied. Numbers/
    dates/None pass through untouched; only str values get the leading
    apostrophe that forces spreadsheet apps to treat the cell as literal text.
    """
    if isinstance(v, str) and v.startswith(_FORMULA_TRIGGER_CHARS):
        return "'" + v
    return v


def render_excel(title: str, columns: list[ColumnSpec], rows: list[dict], header_lines: list[str]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31] or "Export"

    r = 1
    for line in header_lines:
        ws.cell(row=r, column=1, value=_sanitize_cell(line))
        r += 1
    r += 1  # blank separator row

    header_row = r
    for i, c in enumerate(columns, start=1):
        cell = ws.cell(row=header_row, column=i, value=_sanitize_cell(c.label))
        cell.font = Font(bold=True)
    for row in rows:
        r += 1
        for i, c in enumerate(columns, start=1):
            v = row.get(c.key)
            ws.cell(row=r, column=i, value=float(v) if isinstance(v, Decimal) else _sanitize_cell(v))

    for i, c in enumerate(columns, start=1):
        ws.column_dimensions[chr(64 + i) if i <= 26 else "A"].width = max(12, len(c.label) + 2)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
